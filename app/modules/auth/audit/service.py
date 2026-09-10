import csv
import io
import json
import math
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple, Union

from sqlalchemy import desc, func
from sqlalchemy.orm import Session

from app.modules.auth.models import Auditoria, Usuario
from app.modules.auth.audit.schemas import AuditLogEntry, AuditLogListResponse

SENSITIVE_KEYS = {
    "password", "password_hash", "token_version", "refresh_token",
    "token", "access_token", "jwt", "secret"
}


def sanitize_payload(payload: Any) -> Any:
    """Elimina contraseñas y tokens sensibles de los datos a auditar (RN-CU21-02)."""
    if isinstance(payload, dict):
        cleaned = {}
        for k, v in payload.items():
            if k.lower() in SENSITIVE_KEYS:
                cleaned[k] = "[PROTEGIDO]"
            else:
                cleaned[k] = sanitize_payload(v)
        return cleaned
    elif isinstance(payload, list):
        return [sanitize_payload(item) for item in payload]
    elif isinstance(payload, str):
        try:
            parsed = json.loads(payload)
            if isinstance(parsed, (dict, list)):
                return sanitize_payload(parsed)
        except Exception:
            pass
        return payload
    return payload


def registrar_evento(
    db: Session,
    id_usuario: int,
    accion: str,
    id_clinica: Optional[int] = None,
    tabla_afectada: Optional[str] = None,
    registro_id: Optional[int] = None,
    descripcion: Optional[str] = None,
    datos_anteriores: Optional[Any] = None,
    datos_nuevos: Optional[Any] = None,
    direccion_ip: Optional[str] = None,
) -> Auditoria:
    """
    Registra de forma inmutable un evento en la bitácora de auditoría (RN-CU21-01).
    Solo permite inserción: ninguna operación modifica registros existentes.
    """
    if id_clinica is None:
        user = db.query(Usuario).filter(Usuario.id_usuario == id_usuario).first()
        if user and user.id_clinica:
            id_clinica = user.id_clinica
        else:
            id_clinica = 1

    sanitized_ant = sanitize_payload(datos_anteriores) if datos_anteriores else None
    sanitized_nue = sanitize_payload(datos_nuevos) if datos_nuevos else None

    str_ant = (
        json.dumps(sanitized_ant, ensure_ascii=False)
        if isinstance(sanitized_ant, (dict, list))
        else (str(sanitized_ant) if sanitized_ant is not None else None)
    )
    str_nue = (
        json.dumps(sanitized_nue, ensure_ascii=False)
        if isinstance(sanitized_nue, (dict, list))
        else (str(sanitized_nue) if sanitized_nue is not None else None)
    )

    evento = Auditoria(
        id_clinica=id_clinica,
        id_usuario=id_usuario,
        tabla_afectada=tabla_afectada,
        registro_id=registro_id,
        accion=accion.upper(),
        descripcion=descripcion,
        datos_anteriores=str_ant,
        datos_nuevos=str_nue,
        direccion_ip=direccion_ip,
        fecha_hora=datetime.now(),
    )
    db.add(evento)
    db.commit()
    db.refresh(evento)
    return evento


def _build_audit_query(
    db: Session,
    tenant_id: Optional[int],
    fecha_inicio: Optional[datetime] = None,
    fecha_fin: Optional[datetime] = None,
    id_usuario: Optional[int] = None,
    accion: Optional[str] = None,
    tabla_afectada: Optional[str] = None,
    registro_id: Optional[int] = None,
    busqueda: Optional[str] = None,
):
    query = db.query(Auditoria, Usuario).outerjoin(
        Usuario, Auditoria.id_usuario == Usuario.id_usuario
    )

    if tenant_id is not None:
        query = query.filter(Auditoria.id_clinica == tenant_id)

    if fecha_inicio:
        query = query.filter(Auditoria.fecha_hora >= fecha_inicio)
    if fecha_fin:
        query = query.filter(Auditoria.fecha_hora <= fecha_fin)
    if id_usuario:
        query = query.filter(Auditoria.id_usuario == id_usuario)
    if accion:
        query = query.filter(Auditoria.accion.ilike(f"%{accion}%"))
    if tabla_afectada:
        query = query.filter(Auditoria.tabla_afectada.ilike(f"%{tabla_afectada}%"))
    if registro_id:
        query = query.filter(Auditoria.registro_id == registro_id)
    if busqueda:
        pattern = f"%{busqueda}%"
        query = query.filter(
            (Auditoria.descripcion.ilike(pattern))
            | (Auditoria.tabla_afectada.ilike(pattern))
            | (Auditoria.accion.ilike(pattern))
            | (Auditoria.direccion_ip.ilike(pattern))
            | (Usuario.nombres.ilike(pattern))
            | (Usuario.apellidos.ilike(pattern))
            | (Usuario.correo.ilike(pattern))
        )

    return query


def _map_row_to_entry(auditoria: Auditoria, usuario: Optional[Usuario]) -> AuditLogEntry:
    nombre_completo = "Sistema / Desconocido"
    correo_usuario = None
    if usuario:
        nombre_completo = f"{usuario.nombres or ''} {usuario.apellidos or ''}".strip() or usuario.correo
        correo_usuario = usuario.correo

    datos_ant = auditoria.datos_anteriores
    if isinstance(datos_ant, str):
        try:
            datos_ant = json.loads(datos_ant)
        except Exception:
            pass

    datos_nue = auditoria.datos_nuevos
    if isinstance(datos_nue, str):
        try:
            datos_nue = json.loads(datos_nue)
        except Exception:
            pass

    return AuditLogEntry(
        id_auditoria=auditoria.id_auditoria,
        id_clinica=auditoria.id_clinica,
        tenant_id=str(auditoria.id_clinica) if auditoria.id_clinica is not None else None,
        id_usuario=auditoria.id_usuario,
        nombre_usuario=nombre_completo,
        correo_usuario=correo_usuario,
        tabla_afectada=auditoria.tabla_afectada,
        registro_id=auditoria.registro_id,
        accion=auditoria.accion,
        descripcion=auditoria.descripcion,
        datos_anteriores=datos_ant,
        datos_nuevos=datos_nue,
        direccion_ip=auditoria.direccion_ip,
        fecha_hora=auditoria.fecha_hora,
    )


def get_audit_logs(
    db: Session,
    tenant_id: Optional[int],
    page: int = 1,
    page_size: int = 20,
    fecha_inicio: Optional[datetime] = None,
    fecha_fin: Optional[datetime] = None,
    id_usuario: Optional[int] = None,
    accion: Optional[str] = None,
    tabla_afectada: Optional[str] = None,
    registro_id: Optional[int] = None,
    busqueda: Optional[str] = None,
) -> AuditLogListResponse:
    """Consulta paginada con filtros de la bitácora de auditoría (CU21)."""
    query = _build_audit_query(
        db=db,
        tenant_id=tenant_id,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        id_usuario=id_usuario,
        accion=accion,
        tabla_afectada=tabla_afectada,
        registro_id=registro_id,
        busqueda=busqueda,
    )

    total = query.count()
    total_pages = math.ceil(total / page_size) if page_size > 0 else 1

    offset = (page - 1) * page_size
    rows = query.order_by(desc(Auditoria.fecha_hora)).offset(offset).limit(page_size).all()

    items = [_map_row_to_entry(aud, usr) for aud, usr in rows]

    return AuditLogListResponse(
        data=items,
        total=total,
        page=page,
        page_size=page_size,
        total_pages=total_pages,
    )


def get_audit_log_by_id(
    db: Session,
    id_auditoria: int,
    tenant_id: Optional[int],
) -> Optional[AuditLogEntry]:
    """Obtiene el detalle de un registro específico de auditoría."""
    query = db.query(Auditoria, Usuario).outerjoin(
        Usuario, Auditoria.id_usuario == Usuario.id_usuario
    ).filter(Auditoria.id_auditoria == id_auditoria)

    if tenant_id is not None:
        query = query.filter(Auditoria.id_clinica == tenant_id)

    result = query.first()
    if not result:
        return None
    aud, usr = result
    return _map_row_to_entry(aud, usr)


def export_audit_logs_excel(
    db: Session,
    tenant_id: Optional[int],
    fecha_inicio: Optional[datetime] = None,
    fecha_fin: Optional[datetime] = None,
    id_usuario: Optional[int] = None,
    accion: Optional[str] = None,
    tabla_afectada: Optional[str] = None,
    registro_id: Optional[int] = None,
    busqueda: Optional[str] = None,
) -> io.BytesIO:
    """Exporta registros de auditoría filtrados a una hoja de cálculo Excel (.xlsx)."""
    query = _build_audit_query(
        db=db,
        tenant_id=tenant_id,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        id_usuario=id_usuario,
        accion=accion,
        tabla_afectada=tabla_afectada,
        registro_id=registro_id,
        busqueda=busqueda,
    ).order_by(desc(Auditoria.fecha_hora)).limit(1000)

    rows = query.all()

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bitacora Auditoria"

        # Título
        ws.merge_cells("A1:I1")
        ws["A1"] = "Hospital San Juan de Dios - Bitácora de Auditoría (CU21)"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # Subtítulo con fecha
        ws.merge_cells("A2:I2")
        ws["A2"] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Total Registros: {len(rows)}"
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="6B7280")
        ws["A2"].alignment = Alignment(horizontal="center")

        # Cabecera
        headers = ["ID", "Fecha/Hora", "Usuario", "Correo", "Acción", "Tabla Afectada", "ID Registro", "IP", "Descripción"]
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB"),
        )

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[4].height = 24

        row_idx = 5
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        for aud, usr in rows:
            nombre_usr = f"{usr.nombres or ''} {usr.apellidos or ''}".strip() if usr else "Sistema"
            correo_usr = usr.correo if usr else "-"
            fecha_str = aud.fecha_hora.strftime("%d/%m/%Y %H:%M:%S") if aud.fecha_hora else "-"

            data_row = [
                aud.id_auditoria,
                fecha_str,
                nombre_usr,
                correo_usr,
                aud.accion,
                aud.tabla_afectada or "-",
                aud.registro_id or "-",
                aud.direccion_ip or "-",
                aud.descripcion or "-",
            ]

            for col_idx, val in enumerate(data_row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name="Calibri", size=10)
                cell.border = thin_border
                if row_idx % 2 == 0:
                    cell.fill = alt_fill
                if col_idx in [1, 5, 7]:
                    cell.alignment = Alignment(horizontal="center")

            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    except ImportError:
        # Fallback a CSV formateado en BytesIO si openpyxl no estuviera instalado
        output = io.BytesIO()
        output.write("\ufeff".encode("utf-8"))
        text_stream = io.StringIO()
        writer = csv.writer(text_stream, delimiter=";")
        writer.writerow(["ID", "Fecha_Hora", "Usuario", "Correo", "Accion", "Tabla", "ID_Registro", "IP", "Descripcion"])
        for aud, usr in rows:
            nombre_usr = f"{usr.nombres or ''} {usr.apellidos or ''}".strip() if usr else "Sistema"
            writer.writerow([
                aud.id_auditoria,
                aud.fecha_hora.isoformat() if aud.fecha_hora else "",
                nombre_usr,
                usr.correo if usr else "",
                aud.accion,
                aud.tabla_afectada or "",
                aud.registro_id or "",
                aud.direccion_ip or "",
                aud.descripcion or "",
            ])
        output.write(text_stream.getvalue().encode("utf-8"))
        output.seek(0)
        return output


def export_audit_logs_pdf(
    db: Session,
    tenant_id: Optional[int],
    fecha_inicio: Optional[datetime] = None,
    fecha_fin: Optional[datetime] = None,
    id_usuario: Optional[int] = None,
    accion: Optional[str] = None,
    tabla_afectada: Optional[str] = None,
    registro_id: Optional[int] = None,
    busqueda: Optional[str] = None,
) -> io.BytesIO:
    """Genera un archivo PDF formal con los registros de auditoría filtrados."""
    query = _build_audit_query(
        db=db,
        tenant_id=tenant_id,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        id_usuario=id_usuario,
        accion=accion,
        tabla_afectada=tabla_afectada,
        registro_id=registro_id,
        busqueda=busqueda,
    ).order_by(desc(Auditoria.fecha_hora)).limit(500)

    rows = query.all()
    output = io.BytesIO()

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30,
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#1E3A8A"),
            alignment=1,
        )
        subtitle_style = ParagraphStyle(
            "SubtitleStyle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#6B7280"),
            alignment=1,
        )
        cell_style = ParagraphStyle(
            "CellStyle",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#111827"),
        )
        header_cell_style = ParagraphStyle(
            "HeaderCellStyle",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            fontName="Helvetica-Bold",
            textColor=colors.white,
            alignment=1,
        )

        elements.append(Paragraph("Hospital San Juan de Dios - Bitácora de Auditoría", title_style))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(
            f"Reporte de Trazabilidad Clínica y Administrativa | Generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Total Registros: {len(rows)}",
            subtitle_style,
        ))
        elements.append(Spacer(1, 14))

        headers = ["ID", "Fecha/Hora", "Usuario", "Acción", "Tabla", "Reg ID", "IP", "Descripción"]
        data = [[Paragraph(h, header_cell_style) for h in headers]]

        for aud, usr in rows:
            nombre_usr = f"{usr.nombres or ''} {usr.apellidos or ''}".strip() if usr else "Sistema"
            fecha_str = aud.fecha_hora.strftime("%d/%m/%Y %H:%M") if aud.fecha_hora else "-"
            desc_text = (aud.descripcion[:60] + "...") if aud.descripcion and len(aud.descripcion) > 60 else (aud.descripcion or "-")

            data.append([
                Paragraph(str(aud.id_auditoria), cell_style),
                Paragraph(fecha_str, cell_style),
                Paragraph(nombre_usr, cell_style),
                Paragraph(aud.accion, cell_style),
                Paragraph(aud.tabla_afectada or "-", cell_style),
                Paragraph(str(aud.registro_id or "-"), cell_style),
                Paragraph(aud.direccion_ip or "-", cell_style),
                Paragraph(desc_text, cell_style),
            ])

        col_widths = [35, 80, 110, 55, 75, 45, 75, 255]
        t = Table(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ]))
        elements.append(t)
        doc.build(elements)
        output.seek(0)
        return output

    except ImportError:
        # Fallback simple PDF si reportlab no está instalado
        output.write(b"%PDF-1.4\n%Fallback Audit Report\n")
        output.seek(0)
        return output
